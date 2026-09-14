"""
Regression tests for exact text round-tripping through the spreadsheet exporters.

The XLSX exporter already forces every cell to be an explicitly-typed string
(``data_type`` ``"s"``) and the ODS exporter writes every cell with
``office:value-type="string"``, so no additional escaping is needed to keep
formula-like text inert. These tests pin the contract that the stored cell text
equals the original value exactly — in particular that no apostrophe is
prepended to values that merely begin with a formula-like character.
"""

import zipfile
from io import BytesIO
from unittest.mock import MagicMock, patch
from xml.etree import ElementTree

import pytest

TABLE_NS = "urn:oasis:names:tc:opendocument:xmlns:table:1.0"
TEXT_NS = "urn:oasis:names:tc:opendocument:xmlns:text:1.0"
OFFICE_NS = "urn:oasis:names:tc:opendocument:xmlns:office:1.0"

# Every value a spreadsheet application could treat as the start of a formula,
# plus a non-ASCII formula.
DANGEROUS_VALUES = [
    "=CMD",
    "+1",
    "-2",
    "@x",
    "|p",
    "%q",
    "\tTab",
    "\rCR",
    "\nLF",
    "=ورقة",
]


def ods_paragraph_text(value):
    """
    The text an XML parser reports for a value written into ``content.xml``.

    The exporter writes carriage returns as the character reference ``&#13;``,
    so parsers resolve them back to ``\\r`` without applying the XML 1.0
    line-ending normalization that only affects literal CR bytes. The round
    trip must be exact for every character.
    """

    return value


def run_export(table, user, options):
    """
    Run an export job writing into an in-memory file and return the raw bytes.

    This mirrors ``run_export_job_with_mock_storage`` in the CSV export tests but
    returns bytes, because both workbooks are binary.
    """

    storage_mock = MagicMock()
    stub_file = BytesIO()
    storage_mock.open.return_value = stub_file
    stub_file.close = lambda: None

    with patch("jadawel.core.storage.get_default_storage", return_value=storage_mock):
        from jadawel.contrib.database.export.handler import ExportHandler

        handler = ExportHandler()
        job = handler.create_pending_export_job(user, table, None, options)
        handler.run_export_job(job)

    return stub_file.getvalue()


def read_xlsx_rows(payload):
    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(payload), read_only=True)
    worksheet = workbook.active
    return [[cell for cell in row] for row in worksheet.iter_rows(values_only=True)]


def read_ods_rows(payload):
    """Extract the sheet as a list of rows of cell text, straight from the zip."""

    with zipfile.ZipFile(BytesIO(payload)) as archive:
        # S314: parsing our own exporter's generated output in a test, not
        # untrusted user data.
        content = ElementTree.fromstring(archive.read("content.xml"))  # noqa: S314

    rows = []
    for row in content.iter(f"{{{TABLE_NS}}}table-row"):
        values = []
        for cell in row.findall(f"{{{TABLE_NS}}}table-cell"):
            paragraph = cell.find(f"{{{TEXT_NS}}}p")
            values.append(paragraph.text if paragraph is not None else None)
        rows.append(values)
    return rows


def read_ods_cells(payload):
    """Extract the sheet as rows of the raw ``table-cell`` XML elements."""

    with zipfile.ZipFile(BytesIO(payload)) as archive:
        # S314: parsing our own exporter's generated output in a test, not
        # untrusted user data.
        content = ElementTree.fromstring(archive.read("content.xml"))  # noqa: S314

    rows = []
    for row in content.iter(f"{{{TABLE_NS}}}table-row"):
        rows.append(row.findall(f"{{{TABLE_NS}}}table-cell"))
    return rows


@pytest.fixture
def single_value_setup(data_fixture):
    """A table with one text field and a single row holding ``value``."""

    def setup(value):
        user = data_fixture.create_user()
        table = data_fixture.create_database_table(user=user)
        name_field = data_fixture.create_text_field(
            table=table, name="Name", primary=True
        )
        table.get_model().objects.create(**{f"field_{name_field.id}": value})
        return user, table

    return setup


@pytest.mark.django_db
@pytest.mark.parametrize(
    ("exporter_type", "reader", "expected"),
    [
        ("xlsx", read_xlsx_rows, lambda value: value),
        ("ods", read_ods_rows, ods_paragraph_text),
    ],
)
@pytest.mark.parametrize("dangerous", DANGEROUS_VALUES)
def test_spreadsheet_export_stores_formula_like_text_exactly(
    single_value_setup, exporter_type, reader, expected, dangerous
):
    """
    A cell whose text begins with a formula-like character is stored with its
    original text exactly: the explicit string typing alone keeps it inert, so
    no apostrophe escape is added and the round trip is lossless.
    """

    user, table = single_value_setup(dangerous)

    payload = run_export(table, user, {"exporter_type": exporter_type})

    rows = reader(payload)
    assert rows[1][1] == expected(dangerous)


@pytest.mark.django_db
@pytest.mark.parametrize("dangerous", DANGEROUS_VALUES)
def test_xlsx_formula_like_text_is_a_plain_string_cell(single_value_setup, dangerous):
    """The xlsx cell keeps data_type "s" and stores the text without an escape."""

    from openpyxl import load_workbook

    user, table = single_value_setup(dangerous)

    payload = run_export(table, user, {"exporter_type": "xlsx"})

    worksheet = load_workbook(BytesIO(payload)).active
    cell = worksheet.cell(row=2, column=2)
    assert cell.data_type == "s"
    assert "'" not in cell.value
    assert cell.value == dangerous


@pytest.mark.django_db
@pytest.mark.parametrize("dangerous", DANGEROUS_VALUES)
def test_ods_formula_like_text_is_a_string_typed_cell(single_value_setup, dangerous):
    """The ods cell keeps office:value-type="string" and stores the text as-is."""

    user, table = single_value_setup(dangerous)

    payload = run_export(table, user, {"exporter_type": "ods"})

    cell = read_ods_cells(payload)[1][1]
    assert cell.get(f"{{{OFFICE_NS}}}value-type") == "string"
    paragraph = cell.find(f"{{{TEXT_NS}}}p")
    assert paragraph.text == ods_paragraph_text(dangerous)


@pytest.mark.django_db
@pytest.mark.parametrize(
    ("exporter_type", "reader"),
    [("xlsx", read_xlsx_rows), ("ods", read_ods_rows)],
)
@pytest.mark.parametrize("plain", ["hello", "مرحبا", " =SUM(1)"])
def test_spreadsheet_export_plain_text_round_trips(
    single_value_setup, exporter_type, reader, plain
):
    """
    Text that is not stored with a formula-like first character was never
    escaped and must keep round-tripping unchanged. A formula behind leading
    whitespace is protected by the explicit string typing, not by escaping.
    """

    user, table = single_value_setup(plain)

    payload = run_export(table, user, {"exporter_type": exporter_type})

    rows = reader(payload)
    assert rows[1][1] == plain


@pytest.mark.django_db
@pytest.mark.parametrize(
    ("exporter_type", "reader"),
    [("xlsx", read_xlsx_rows), ("ods", read_ods_rows)],
)
@pytest.mark.parametrize(
    "newlines",
    ["a\rb", "a\nb", "a\r\nb", "\r\nstart", "end\r", "\r\n\r\n"],
)
def test_spreadsheet_export_preserves_newlines_exactly(
    single_value_setup, exporter_type, reader, newlines
):
    """
    CR, LF and CRLF all survive an export/import round trip byte for byte.
    The ODS writer must emit CR as ``&#13;`` because a literal CR byte in XML
    character data is normalized to LF by every conforming parser.
    """

    user, table = single_value_setup(newlines)

    payload = run_export(table, user, {"exporter_type": exporter_type})

    rows = reader(payload)
    assert rows[1][1] == newlines


@pytest.mark.django_db
def test_ods_writes_cr_as_character_reference(single_value_setup):
    """
    The ODS package must not contain a literal CR byte in character data:
    the character reference is what keeps parsers from normalizing it.
    """

    user, table = single_value_setup("a\rb")

    payload = run_export(table, user, {"exporter_type": "ods"})

    with zipfile.ZipFile(BytesIO(payload)) as archive:
        raw = archive.read("content.xml")
    # S314: parsing our own exporter's generated output in a test, not
    # untrusted user data.
    content = ElementTree.fromstring(raw)  # noqa: S314
    data_row = list(content.iter(f"{{{TABLE_NS}}}table-row"))[1]
    paragraph = data_row.findall(f"{{{TABLE_NS}}}table-cell")[1].find(f"{{{TEXT_NS}}}p")
    assert paragraph.text == "a\rb"
    assert b"a&#13;b" in raw
    # No literal CR byte may sit in character data.
    assert b"\r" not in raw.replace(b"a&#13;b", b"")
